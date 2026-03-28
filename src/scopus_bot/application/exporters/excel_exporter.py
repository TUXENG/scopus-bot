from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from scopus_bot.data.models import Document


class ExcelExporter:
    HEADERS = [
        "Title",
        "Document Type",
        "Authors",
        "Source",
        "Year",
        "Citations",
        "DOI / Link",
    ]

    def export_documents_by_keyword(
        self,
        documents_by_keyword: dict[str, list[Document]],
        output_path: Path,
    ) -> Path:
        workbook = Workbook()

        default_sheet = workbook.active
        workbook.remove(default_sheet)

        for keyword, documents in documents_by_keyword.items():
            sheet_title = self._build_sheet_title(keyword)
            sheet = workbook.create_sheet(title=sheet_title)

            sheet.append(self.HEADERS)

            for cell in sheet[1]:
                cell.font = Font(bold=True)

            for document in documents:
                sheet.append(
                    [
                        document.title,
                        document.doc_type,
                        document.authors,
                        document.source,
                        document.year,
                        document.citations,
                        document.doi,
                    ]
                )

            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            self._adjust_column_widths(sheet)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)

        return output_path

    def _build_sheet_title(self, keyword: str) -> str:
        invalid_chars = ['\\', '/', '*', '?', ':', '[', ']']
        title = keyword.strip() or "Sin keyword"

        for char in invalid_chars:
            title = title.replace(char, "_")

        return title[:31]

    def _adjust_column_widths(self, sheet) -> None:
        for column_cells in sheet.columns:
            max_length = 0
            column_index = column_cells[0].column
            column_letter = get_column_letter(column_index)

            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                if len(value) > max_length:
                    max_length = len(value)

            sheet.column_dimensions[column_letter].width = min(max_length + 2, 60)